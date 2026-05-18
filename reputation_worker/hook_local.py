from __future__ import annotations

import os
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


@dataclass(slots=True)
class HookClientLocal:
    """Salva lotes de registros em uma planilha .xlsx local."""

    file_path: str
    sheet_name: str = "dados"
    _headers_written: bool = field(default=False, init=False, repr=False)

    def __post_init__(self) -> None:
        self._headers_written = os.path.isfile(self.file_path)

    def post(self, payload: dict[str, Any]) -> dict[str, Any]:
        records: list[dict[str, Any]] = payload.get("records", [])
        if not records:
            return {"status_code": 204, "rows_written": 0}

        if os.path.isfile(self.file_path):
            wb = load_workbook(self.file_path)
            ws: Worksheet = (
                wb[self.sheet_name]
                if self.sheet_name in wb.sheetnames
                else wb.create_sheet(self.sheet_name)
            )
        else:
            wb = Workbook()
            ws = wb.active  # type: ignore[assignment]
            ws.title = self.sheet_name
            self._headers_written = False

        headers = list(records[0].keys())

        if not self._headers_written or ws.max_row == 0:
            ws.append(headers)
            self._headers_written = True

        for record in records:
            ws.append([record.get(h, "") for h in headers])

        wb.save(self.file_path)

        return {"status_code": 201, "rows_written": len(records)}
