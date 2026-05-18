import os
import tempfile
import unittest

from openpyxl import load_workbook

from reputation_worker.hook_local import HookClientLocal


class TestHookClientLocal(unittest.TestCase):
    def _tmp_xlsx(self) -> str:
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        os.unlink(path)
        return path

    def test_creates_file_and_writes_headers_and_rows(self):
        path = self._tmp_xlsx()
        try:
            hook = HookClientLocal(file_path=path, sheet_name="dados")
            result = hook.post({"records": [{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]})

            self.assertEqual(result["rows_written"], 2)
            self.assertTrue(os.path.isfile(path))

            wb = load_workbook(path)
            ws = wb["dados"]
            rows = list(ws.values)

            self.assertEqual(rows[0], ("a", "b"))  # cabecalho
            self.assertEqual(rows[1], (1, "x"))
            self.assertEqual(rows[2], (2, "y"))
        finally:
            if os.path.isfile(path):
                os.unlink(path)

    def test_appends_rows_on_second_call(self):
        path = self._tmp_xlsx()
        try:
            hook = HookClientLocal(file_path=path, sheet_name="dados")
            hook.post({"records": [{"a": 1, "b": "x"}]})
            hook2 = HookClientLocal(file_path=path, sheet_name="dados")
            hook2.post({"records": [{"a": 2, "b": "y"}]})

            wb = load_workbook(path)
            ws = wb["dados"]
            rows = list(ws.values)

            self.assertEqual(rows[0], ("a", "b"))  # cabecalho aparece uma vez
            self.assertEqual(rows[1], (1, "x"))
            self.assertEqual(rows[2], (2, "y"))
        finally:
            if os.path.isfile(path):
                os.unlink(path)

    def test_empty_records_returns_204(self):
        path = self._tmp_xlsx()
        try:
            hook = HookClientLocal(file_path=path)
            result = hook.post({"records": []})
            self.assertEqual(result["status_code"], 204)
            self.assertFalse(os.path.isfile(path))
        finally:
            if os.path.isfile(path):
                os.unlink(path)


if __name__ == "__main__":
    unittest.main()
